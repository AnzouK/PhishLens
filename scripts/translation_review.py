#!/usr/bin/env python3
"""Hand a translation to a native speaker, then bring the corrections back.

    python3 scripts/translation_review.py export ha -o hausa-review.xlsx
    python3 scripts/translation_review.py import ha hausa-review.xlsx

export: one workbook with the English text and the current translation of
every extension and website string, and an empty Correction column.
import: writes each filled Correction into extension/lib/locales/<lang>.js
and site-src/strings/<lang>.json; then run the locale tests and
scripts/build_site.py. Needs openpyxl.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
LOCALES = ROOT / "extension" / "lib" / "locales"
STRINGS = ROOT / "site-src" / "strings"
LINE = re.compile(r'^(\s*)"([a-z0-9_]+)":\s*(".*")(,?)\s*$')
HEADERS = ["Key", "Where it appears", "English", "Current translation", "Correction", "Comment"]
WHERE = [
    (("flag_", "reason_"), "Attachment warnings"),
    (("att_", "drop_", "file_"), "Attachment results"),
    (
        ("chip_", "badge_", "threat_", "auth_", "dkim_", "links_", "url_"),
        "Signal labels in a verdict",
    ),
    (("fb_", "share_", "review_", "live_", "stat_", "chart_", "lang_"), "Reports and Insights"),
    (("auto_", "row_", "notify_"), "Automatic scan in Gmail"),
    (("settings_", "custom_", "credits_", "about_", "connected", "testing"), "Settings"),
    (
        ("verdict_", "sub_", "lead_", "meter_", "agent_", "why_", "token_", "scan_"),
        "Verdict banner and popup",
    ),
]


def extension_strings(lang: str) -> dict[str, str]:
    out = {}
    for line in (LOCALES / f"{lang}.js").read_text(encoding="utf-8").splitlines():
        m = LINE.match(line)
        if m:
            out[m.group(2)] = json.loads(m.group(3))
    return out


def where(key: str) -> str:
    return next(
        (label for prefixes, label in WHERE if key.startswith(prefixes)), "Popup and banners"
    )


def site_where(key: str) -> str:
    if key.startswith("js_"):
        return "Website: messages of the Try it box"
    if key.startswith(("legal_", "privacy_", "terms_")):
        return "Website: legal pages"
    return "Website: home page"


def export(lang: str, out: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook()
    info = wb.active
    info.title = "Read me"
    rows = [
        [f"PhishLens: review of the {lang} translation"],
        [],
        ["What to do"],
        [
            "Read the Current translation column. If a text is wrong, unclear or unnatural, write the "
            "better text in the yellow Correction column. Leave Correction empty when the text is fine."
        ],
        [
            "Keep everything written {like_this} exactly as it is: the program puts a name, a number or a "
            "file name there. Keep HTML tags such as <b> and </b> around the same words."
        ],
        ["Use Comment for questions or remarks. Only the Correction column is copied back."],
        [],
        ["Example (Extension tab)"],
        [
            "Key: verdict_phishing · English: This email looks like phishing · Current: (the translation) · "
            "Correction: the better wording, or empty"
        ],
        [],
        ["Tabs: Extension (the Chrome extension) and Website (phishlens.dev)."],
    ]
    for r in rows:
        info.append(r)
    for row in info.iter_rows():
        for c in row:
            c.font, c.alignment = font, Alignment(wrap_text=True, vertical="top")
    for r in (1, 3, 8):
        info.cell(r, 1).font = Font(name="Arial", size=12 if r == 1 else 10, bold=True)
    info.column_dimensions["A"].width = 110

    sources = {
        "Extension": (extension_strings("en"), extension_strings(lang), where),
        "Website": (
            json.loads((STRINGS / "en.json").read_text(encoding="utf-8")),
            json.loads((STRINGS / f"{lang}.json").read_text(encoding="utf-8")),
            site_where,
        ),
    }
    for title, (en, tr, place) in sources.items():
        ws = wb.create_sheet(title)
        ws.append(HEADERS)
        for key, text in en.items():
            ws.append([key, place(key), text, tr.get(key, ""), "", ""])
        for row in ws.iter_rows():
            for c in row:
                c.font, c.alignment = font, wrap
        for c in ws[1]:
            c.font = bold
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 5).fill = fill
        for col, width in zip("ABCDEF", (22, 24, 48, 48, 48, 30), strict=True):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    print(f"{out}: {sum(len(s[0]) for s in sources.values())} texts to review")


def _apply_extension(lang: str, fixes: dict[str, str]) -> int:
    path = LOCALES / f"{lang}.js"
    lines, n = [], 0
    for line in path.read_text(encoding="utf-8").splitlines():
        m = LINE.match(line)
        if m and m.group(2) in fixes:
            line = f'{m.group(1)}"{m.group(2)}": {json.dumps(fixes[m.group(2)], ensure_ascii=False)}{m.group(4)}'
            n += 1
        lines.append(line)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return n


def _apply_site(lang: str, fixes: dict[str, str]) -> int:
    path = STRINGS / f"{lang}.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    n = sum(1 for k in fixes if k in data)
    data.update({k: v for k, v in fixes.items() if k in data})
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return n


def import_(lang: str, workbook: str) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(workbook, read_only=True)
    done = {}
    for title, apply in (("Extension", _apply_extension), ("Website", _apply_site)):
        fixes = {}
        for row in wb[title].iter_rows(min_row=2, values_only=True):
            key, fix = row[0], row[4]
            if key and fix is not None and str(fix).strip():
                fixes[str(key)] = str(fix).strip()
        done[title] = apply(lang, fixes)
    print(f"corrections copied: extension {done['Extension']}, website {done['Website']}")
    print("next: node --test tests/extension/ && python3 scripts/build_site.py")


def main() -> None:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("lang")
    e.add_argument("-o", "--out", required=True)
    i = sub.add_parser("import")
    i.add_argument("lang")
    i.add_argument("workbook")
    a = ap.parse_args()
    export(a.lang, a.out) if a.cmd == "export" else import_(a.lang, a.workbook)


if __name__ == "__main__":
    main()
